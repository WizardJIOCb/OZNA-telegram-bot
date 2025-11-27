# -*- coding: utf-8 -*-
"""
CSV to Excel Converter for Bot Reports
Converts bot-generated CSV reports into styled Excel files
"""
import csv
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo
from datetime import datetime

# -------------------------------------------------
# 1. Path to your CSV file (change if needed)
csv_path = "simple_report_all_2025-10-25_17-49-05.csv"

# 2. Read CSV file (comma separator, double quotes)
df = pd.read_csv(csv_path, sep=',', quotechar='"', encoding='utf-8-sig',
                 dtype=str, keep_default_na=False)

# 3. Column names are already in the CSV file (first row)
#    If they are missing, uncomment and specify:
# df.columns = ['Проект','Дата','Процесс','Вид деятельности','Основная активность',
#               'Тип работ 1','Тип работ 2','Тип работ 3','ФИО']

# -------------------------------------------------
# 4. Format dates (if in dd.mm.yyyy format)
if 'Дата' in df.columns:
    df['Дата'] = pd.to_datetime(df['Дата'], format='%d.%m.%Y', errors='coerce') \
                  .dt.strftime('%d.%m.%Y')

# 5. Sort by date and employee name
if {'Дата', 'ФИО'}.issubset(df.columns):
    df['Дата_sort'] = pd.to_datetime(df['Дата'], format='%d.%m.%Y', errors='coerce')
    df = df.sort_values(['Дата_sort', 'ФИО']).drop(columns='Дата_sort')
    df.insert(df.columns.get_loc('Дата') + 1, 'Дата', df.pop('Дата'))

# -------------------------------------------------
# 6. Create .xlsx file
wb = Workbook()
ws = wb.active
ws.title = "Табель"

# ---- Headers -------------------------------------------------
headers = list(df.columns)
for c, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1f4e79")  # Dark blue
    cell.alignment = Alignment(horizontal="center", vertical="center")

# ---- Data ----------------------------------------------------
for r, row in enumerate(dataframe_to_rows(df, index=False, header=False), start=2):
    for c, val in enumerate(row, 1):
        cell = ws.cell(row=r, column=c, value=val)
        # Red text for idle time (Простой)
        if isinstance(val, str) and 'Простой' in val:
            cell.font = Font(color="FF0000", bold=True)

# ---- Table with filters ---------------------------------------
tab = Table(displayName="Tab", ref=f"A1:{chr(64+len(headers))}{len(df)+1}")
style = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
tab.tableStyleInfo = style
ws.add_table(tab)

# ---- Auto-width columns ---------------------------------------
for col in ws.columns:
    max_len = 0
    column = col[0].column_letter
    for cell in col:
        try:
            if len(str(cell.value)) > max_len:
                max_len = len(str(cell.value))
        except:
            pass
    ws.column_dimensions[column].width = min(max_len + 2, 50)

# -------------------------------------------------
# 7. Save the file
out_name = f"report_{datetime.now():%Y-%m-%d_%H-%M-%S}.xlsx"
wb.save(out_name)
print(f"✅ Done! File created: {out_name}")
